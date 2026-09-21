"""KLASSIO — l'établissement reprend ses données.

Un export sort les données hors de tout contrôle d'accès : c'est l'endroit où
une erreur d'isolation coûte le plus cher. Ces tests éprouvent d'abord les
refus, ensuite le contenu.
"""
import io
import json
import os
import sys
import unittest
import zipfile

BACKEND = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, BACKEND)

TEST_DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "klassio_export_test.db")
os.environ["KLASSIO_DB_PATH"] = TEST_DB

import db  # noqa: E402
db.DB_PATH = os.path.abspath(TEST_DB)

import app as flask_app_module  # noqa: E402
import security  # noqa: E402
import exports as exports_module  # noqa: E402


def setUpModule():
    if os.path.exists(db.DB_PATH):
        os.remove(db.DB_PATH)
    db.init_db()


def tearDownModule():
    if os.path.exists(db.DB_PATH):
        os.remove(db.DB_PATH)


class Base(unittest.TestCase):
    def setUp(self):
        self.c = flask_app_module.app.test_client()
        security.reset_rate_limits_for_tests()

    def ecole(self, suffixe, eleves=("Kabeya", "Mbuyi")):
        r = self.c.post("/api/auth/register-school", json={
            "email": f"dir.{suffixe}@exp.test", "password": "Secret123!",
            "name": "Directeur Export", "school_name": f"École {suffixe}"})
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        d = r.get_json()
        h = {"Authorization": f"Bearer {d['token']}"}
        an = self.c.get("/api/academic-years", headers=h).get_json()[0]["id"]
        cls = self.c.post("/api/classes", json={"academic_year_id": an, "name": "6e A"},
                          headers=h).get_json()
        ids = []
        for nom in eleves:
            ids.append(self.c.post("/api/students", json={
                "academic_year_id": an, "class_id": cls["id"],
                "first_name": "Jean", "last_name": nom}, headers=h).get_json()["id"])
        return {"h": h, "tenant_id": d["tenant_id"], "year": an, "class": cls, "students": ids}

    def archive(self, h, export_id):
        r = self.c.get(f"/api/exports/{export_id}/download", headers=h)
        # Le corps est un ZIP : on ne le décode en texte que s'il a échoué,
        # sinon la composition du message d'erreur plante sur du binaire.
        if r.status_code != 200:
            self.fail(f"téléchargement refusé ({r.status_code}) : {r.get_data(as_text=True)[:200]}")
        return zipfile.ZipFile(io.BytesIO(r.data))

    def noms_dans_feuille(self, z, feuille):
        """Les valeurs réellement écrites dans un classeur de l'archive."""
        import openpyxl
        nom = [n for n in z.namelist() if n.endswith(f"{feuille}.xlsx")]
        if not nom:
            return []
        wb = openpyxl.load_workbook(io.BytesIO(z.read(nom[0])), read_only=True)
        valeurs = []
        for ligne in wb[wb.sheetnames[0]].iter_rows(values_only=True):
            valeurs.extend(str(v) for v in ligne if v is not None)
        return valeurs


class ExportsFonctionnels(Base):
    def test_01_un_export_annuel_produit_une_archive_reelle(self):
        ctx = self.ecole("fonc01")
        r = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": ctx["year"]},
                        headers=ctx["h"])
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        corps = r.get_json()
        self.assertEqual(corps["status"], "READY")
        self.assertGreater(corps["file_size"], 0)
        self.assertEqual(corps["counts"].get("Eleves"), 2)

        z = self.archive(ctx["h"], corps["id"])
        noms = z.namelist()
        self.assertTrue(any(n.endswith("Manifest.json") for n in noms), noms)
        self.assertTrue(any(n.endswith("Eleves.xlsx") for n in noms), noms)

    def test_02_le_manifeste_dit_la_verite_et_ne_fuite_rien(self):
        ctx = self.ecole("fonc02")
        eid = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": ctx["year"]},
                          headers=ctx["h"]).get_json()["id"]
        z = self.archive(ctx["h"], eid)
        nom = [n for n in z.namelist() if n.endswith("Manifest.json")][0]
        m = json.loads(z.read(nom).decode("utf-8"))
        self.assertEqual(m["feuilles"]["Eleves"], 2, "le manifeste annonce un compte faux")
        self.assertIn("École fonc02", m["etablissement"])
        brut = json.dumps(m).lower()
        for secret in ("password", "api_key", "token", "secret", "/users/"):
            self.assertNotIn(secret, brut, f"« {secret} » ne doit pas figurer dans un manifeste")

    def test_03_aucun_fichier_vide_n_est_ajoute(self):
        """Un classeur vide dans l'archive laisse croire à une perte de
        données. La vérité est dans le manifeste."""
        ctx = self.ecole("fonc03")
        eid = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": ctx["year"]},
                          headers=ctx["h"]).get_json()["id"]
        z = self.archive(ctx["h"], eid)
        self.assertFalse(any(n.endswith("Presences.xlsx") for n in z.namelist()),
                         "une feuille sans aucune ligne a été ajoutée")
        m = json.loads(z.read([n for n in z.namelist() if n.endswith("Manifest.json")][0]))
        self.assertEqual(m["feuilles"]["Presences"], 0, "le manifeste doit quand même la mentionner")

    def test_04_un_export_cible_ne_contient_que_son_jeu(self):
        ctx = self.ecole("fonc04")
        eid = self.c.post("/api/exports", json={"kind": "students", "academic_year_id": ctx["year"]},
                          headers=ctx["h"]).get_json()["id"]
        z = self.archive(ctx["h"], eid)
        feuilles = [n for n in z.namelist() if n.endswith(".xlsx")]
        self.assertEqual(len(feuilles), 1)
        self.assertTrue(feuilles[0].endswith("Eleves.xlsx"))

    def test_05_l_historique_garde_la_trace(self):
        ctx = self.ecole("fonc05")
        self.c.post("/api/exports", json={"kind": "students"}, headers=ctx["h"])
        liste = self.c.get("/api/exports", headers=ctx["h"]).get_json()
        self.assertEqual(len(liste), 1)
        self.assertEqual(liste[0]["status"], "READY")
        self.assertEqual(liste[0]["requested_by_name"], "Directeur Export")


class IsolationDesExports(Base):
    def test_10_une_ecole_n_exporte_jamais_les_eleves_d_une_autre(self):
        """Le test central. Deux écoles, des noms d'élèves distincts : aucun ne
        doit traverser."""
        a = self.ecole("iso10a", eleves=("AlphaUnique",))
        b = self.ecole("iso10b", eleves=("BetaUnique",))
        eid = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": a["year"]},
                          headers=a["h"]).get_json()["id"]
        z = self.archive(a["h"], eid)
        # On LIT le classeur au lieu de fouiller des octets : un .xlsx est
        # lui-même une archive compressée, y chercher une chaîne ne prouve rien.
        noms = self.noms_dans_feuille(z, "Eleves")
        self.assertIn("AlphaUnique", noms, f"l'élève de l'école A manque : {noms}")
        self.assertNotIn("BetaUnique", noms,
                         "un élève d'une autre école figure dans l'archive")

    def test_11_l_identifiant_d_un_export_d_une_autre_ecole_ne_donne_rien(self):
        """IDOR : posséder l'identifiant ne vaut pas autorisation."""
        a = self.ecole("iso11a")
        b = self.ecole("iso11b")
        eid = self.c.post("/api/exports", json={"kind": "students"}, headers=a["h"]).get_json()["id"]
        r = self.c.get(f"/api/exports/{eid}/download", headers=b["h"])
        self.assertEqual(r.status_code, 404)

    def test_12_une_annee_d_une_autre_ecole_est_refusee(self):
        a = self.ecole("iso12a")
        b = self.ecole("iso12b")
        r = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": b["year"]},
                        headers=a["h"])
        self.assertEqual(r.status_code, 404,
                         "un identifiant d'année d'une autre école a été accepté")

    def test_13_exporter_n_est_pas_lire(self):
        """Un professeur consulte sa classe. Il ne sort pas un fichier de
        l'établissement — c'est un autre acte, et une autre permission."""
        ctx = self.ecole("iso13")
        inv = self.c.post("/api/invitations", json={
            "role": "professeur", "class_ids": [ctx["class"]["id"]]}, headers=ctx["h"]).get_json()
        prof = self.c.post("/api/invitations/accept", json={
            "token": inv["token"], "name": "Prof", "email": "prof13@exp.test",
            "password": "Secret123!"}).get_json()
        h_prof = {"Authorization": f"Bearer {prof['token']}"}
        self.assertEqual(self.c.post("/api/exports", json={"kind": "students"},
                                     headers=h_prof).status_code, 403)
        self.assertEqual(self.c.get("/api/exports", headers=h_prof).status_code, 403)

    def test_14_un_parent_n_exporte_rien(self):
        ctx = self.ecole("iso14")
        inv = self.c.post("/api/invitations", json={
            "role": "parent", "student_ids": [ctx["students"][0]]}, headers=ctx["h"]).get_json()
        par = self.c.post("/api/invitations/accept", json={
            "token": inv["token"], "name": "Parent", "email": "parent14@exp.test",
            "password": "Secret123!"}).get_json()
        h = {"Authorization": f"Bearer {par['token']}"}
        self.assertEqual(self.c.post("/api/exports", json={"kind": "students"}, headers=h).status_code, 403)

    def test_15_sans_session_aucun_export(self):
        ctx = self.ecole("iso15")
        eid = self.c.post("/api/exports", json={"kind": "students"}, headers=ctx["h"]).get_json()["id"]
        self.assertEqual(self.c.get(f"/api/exports/{eid}/download").status_code, 401)
        self.assertEqual(self.c.post("/api/exports", json={"kind": "students"}).status_code, 401)


class IntegriteDesDonnees(Base):
    def test_20_les_montants_exportes_correspondent_a_la_base(self):
        ctx = self.ecole("integ20")
        an, h = ctx["year"], ctx["h"]
        art = self.c.post("/api/catalog-items", json={
            "name": "Frais", "amount": 250, "currency": "USD"}, headers=h).get_json()
        ob = self.c.post("/api/obligations", json={
            "student_id": ctx["students"][0], "academic_year_id": an,
            "catalog_item_id": art["id"]}, headers=h).get_json()
        self.c.post("/api/payments", json={
            "obligation_id": ob["id"], "amount": 100, "method": "cash",
            "idempotency_key": "integ20-1"}, headers=h)

        eid = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": an},
                          headers=h).get_json()["id"]
        z = self.archive(h, eid)
        m = json.loads(z.read([n for n in z.namelist() if n.endswith("Manifest.json")][0]))
        self.assertEqual(m["feuilles"]["Obligations"], 1)
        self.assertEqual(m["feuilles"]["Paiements"], 1)
        self.assertEqual(m["feuilles"]["Recus"], 1, "le reçu du paiement confirmé manque")

    def test_21_une_intention_non_confirmee_n_est_pas_un_paiement(self):
        """Un mobile_money CREATED n'est pas de l'argent. Le faire figurer
        dans une archive comptable serait inexact."""
        ctx = self.ecole("integ21")
        an, h = ctx["year"], ctx["h"]
        art = self.c.post("/api/catalog-items", json={"name": "F", "amount": 90}, headers=h).get_json()
        ob = self.c.post("/api/obligations", json={
            "student_id": ctx["students"][0], "academic_year_id": an,
            "catalog_item_id": art["id"]}, headers=h).get_json()
        self.c.post("/api/payments", json={
            "obligation_id": ob["id"], "amount": 90, "method": "mobile_money",
            "idempotency_key": "integ21-1"}, headers=h)
        eid = self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": an},
                          headers=h).get_json()["id"]
        z = self.archive(h, eid)
        m = json.loads(z.read([n for n in z.namelist() if n.endswith("Manifest.json")][0]))
        self.assertEqual(m["feuilles"]["Paiements"], 0,
                         "une intention non confirmée a été exportée comme un paiement")

    def test_22_exporter_ne_detruit_rien(self):
        """Archiver est une copie, jamais une suppression déguisée."""
        ctx = self.ecole("integ22")
        avant = len(self.c.get("/api/students", headers=ctx["h"]).get_json())
        self.c.post("/api/exports", json={"kind": "annual", "academic_year_id": ctx["year"]},
                    headers=ctx["h"])
        apres = len(self.c.get("/api/students", headers=ctx["h"]).get_json())
        self.assertEqual(avant, apres)
        self.assertEqual(avant, 2)


if __name__ == "__main__":
    unittest.main()
